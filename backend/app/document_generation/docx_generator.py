"""
Generate Word (.docx) documents from resume/cover letter data using python-docx.
Each visual template maps to a style configuration applied consistently.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


TEMPLATES = {
    "classic": {"name_size": 20, "section_size": 12, "body_size": 10, "accent_color": RGBColor(0x1a, 0x1a, 0x2e)},
    "modern":  {"name_size": 22, "section_size": 11, "body_size": 10, "accent_color": RGBColor(0x0f, 0x3d, 0x66)},
    "minimal": {"name_size": 18, "section_size": 11, "body_size": 10, "accent_color": RGBColor(0x33, 0x33, 0x33)},
}


def generate_docx(resume, experiences: list, template_name: str, output_path: str) -> None:
    """Generate a resume .docx file."""
    cfg = TEMPLATES.get(template_name, TEMPLATES["classic"])
    doc = Document()

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    profile = resume.profile

    # ── Header: Name ────────────────────────────────────────
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(profile.full_name)
    name_run.bold = True
    name_run.font.size = Pt(cfg["name_size"])
    name_run.font.color.rgb = cfg["accent_color"]

    # Contact line
    contact_parts = [x for x in [profile.email, profile.phone, profile.location, profile.linkedin_url] if x]
    if contact_parts:
        contact_para = doc.add_paragraph(" | ".join(contact_parts))
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in contact_para.runs:
            run.font.size = Pt(9)

    _add_horizontal_rule(doc)

    sections_order = resume.sections_order or ["summary", "experience", "skills", "education"]
    sections_config = resume.sections_config or {}
    exp_map = {e.id: e for e in experiences}
    ordered_exp = [exp_map[eid] for eid in (resume.selected_experience_ids or []) if eid in exp_map]

    for section_key in sections_order:
        sec_cfg = sections_config.get(section_key, {})
        if not sec_cfg.get("visible", True):
            continue

        if section_key == "summary" and resume.custom_summary:
            _add_section_heading(doc, sec_cfg.get("title", "Professional Summary"), cfg)
            doc.add_paragraph(resume.custom_summary)

        elif section_key == "experience" and ordered_exp:
            _add_section_heading(doc, sec_cfg.get("title", "Experience"), cfg)
            for exp in ordered_exp:
                _add_experience_entry(doc, exp, cfg)

        elif section_key == "skills":
            skills_data = resume.skills_selection or {}
            selected_skills = skills_data.get("selected", [])
            if selected_skills:
                _add_section_heading(doc, sec_cfg.get("title", "Skills"), cfg)
                doc.add_paragraph(", ".join(selected_skills))

        elif section_key == "education":
            # Education stored in profile — placeholder for now
            pass

    doc.save(output_path)


def _add_section_heading(doc: Document, title: str, cfg: dict) -> None:
    para = doc.add_paragraph()
    run = para.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(cfg["section_size"])
    run.font.color.rgb = cfg["accent_color"]
    _add_paragraph_border_bottom(para)


def _add_experience_entry(doc: Document, exp, cfg: dict) -> None:
    # Company + Title row
    title_para = doc.add_paragraph()
    company_run = title_para.add_run(exp.company_name)
    company_run.bold = True
    company_run.font.size = Pt(cfg["body_size"] + 1)

    title_para.add_run(f"  |  {exp.job_title}")

    # Date + location
    date_parts = []
    if exp.start_date:
        date_parts.append(exp.start_date)
    date_parts.append(exp.end_date if exp.end_date else "Present")
    date_str = " – ".join(date_parts)
    if exp.location:
        date_str += f"  •  {exp.location}"

    date_para = doc.add_paragraph(date_str)
    for run in date_para.runs:
        run.italic = True
        run.font.size = Pt(cfg["body_size"] - 1)

    # Bullets
    for bullet in (exp.bullets or []):
        text = bullet.get("text", "") if isinstance(bullet, dict) else str(bullet)
        if text:
            bullet_para = doc.add_paragraph(style="List Bullet")
            bullet_para.add_run(text).font.size = Pt(cfg["body_size"])

    doc.add_paragraph()  # spacer


def _add_horizontal_rule(doc: Document) -> None:
    para = doc.add_paragraph()
    _add_paragraph_border_bottom(para)


def _add_paragraph_border_bottom(para) -> None:
    p = para._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    pBdr.append(bottom)
    pPr.append(pBdr)


def generate_cover_letter_docx(cover_letter, output_path: str) -> None:
    """Generate a cover letter .docx file."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(3.0)

    job = cover_letter.job_application
    if job:
        profile = job.profile
        if profile:
            header = doc.add_paragraph(profile.full_name)
            header.runs[0].bold = True
            header.runs[0].font.size = Pt(14)

        doc.add_paragraph(f"{job.company_name}\n{job.job_title}")
        doc.add_paragraph()

    # Body paragraphs
    for para_text in cover_letter.body_plain.split("\n\n"):
        if para_text.strip():
            doc.add_paragraph(para_text.strip())

    doc.save(output_path)
